"""Excel test data reader using openpyxl.

Parses .xlsx worksheets into normalized lists of dictionaries suitable for
pytest parameterization.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Union

import openpyxl

from utilities.logger import get_logger

logger = get_logger(__name__)


class ExcelReader:
    """Reads structured test data from Excel (.xlsx) files."""

    def __init__(self, file_path: Union[str, Path]) -> None:
        self.file_path = Path(file_path)
        if not self.file_path.exists():
            raise FileNotFoundError(f"Excel test data file not found: {self.file_path}")

    def read_sheet(self, sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
        """Reads rows from the specified sheet (or active sheet) and returns list of dicts.

        Header row (row 1) determines dictionary keys. Empty rows are skipped.
        """
        logger.info("Reading Excel test data from: %s (Sheet: %s)", self.file_path, sheet_name or "active")
        workbook = openpyxl.load_workbook(self.file_path, data_only=True)

        sheet = workbook[sheet_name] if sheet_name and sheet_name in workbook.sheetnames else workbook.active
        if sheet is None:
            raise ValueError(f"Could not open valid worksheet in {self.file_path}")

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            logger.warning("Excel sheet '%s' in %s is completely empty.", sheet.title, self.file_path)
            return []

        # Row 1 is header
        raw_headers = rows[0]
        headers: List[str] = [str(h).strip() if h is not None else f"col_{idx}" for idx, h in enumerate(raw_headers)]

        data_rows: List[Dict[str, Any]] = []
        for row_idx, row in enumerate(rows[1:], start=2):
            # Skip completely empty rows
            if not any(cell is not None and str(cell).strip() != "" for cell in row):
                continue

            row_dict: Dict[str, Any] = {}
            for col_idx, header in enumerate(headers):
                cell_value = row[col_idx] if col_idx < len(row) else None
                # Normalize string representations
                if isinstance(cell_value, str):
                    cell_value = cell_value.strip()
                row_dict[header] = cell_value

            data_rows.append(row_dict)

        logger.info("Successfully loaded %d records from Excel file: %s", len(data_rows), self.file_path.name)
        workbook.close()
        return data_rows


def read_excel_data(file_path: Union[str, Path], sheet_name: Optional[str] = None) -> List[Dict[str, Any]]:
    """Convenience function to read test data from an Excel file."""
    reader = ExcelReader(file_path)
    return reader.read_sheet(sheet_name)
