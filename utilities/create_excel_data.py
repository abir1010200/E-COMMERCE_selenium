import json
import sys
from pathlib import Path

# Ensure workspace root is in sys.path
WORKSPACE_ROOT = Path(__file__).resolve().parent.parent
if str(WORKSPACE_ROOT) not in sys.path:
    sys.path.insert(0, str(WORKSPACE_ROOT))

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config.config import EXCEL_DATA_PATH, JSON_DATA_PATH, TEST_DATA_DIR


def generate_excel_data() -> Path:
    """Reads test_data.json and generates a professionally formatted test_data.xlsx."""
    TEST_DATA_DIR.mkdir(parents=True, exist_ok=True)

    with open(JSON_DATA_PATH, "r", encoding="utf-8") as f:
        records = json.load(f)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PurchaseTests"

    # Header styling
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data styling
    data_font = Font(name="Calibri", size=10)
    data_alignment = Alignment(vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    alt_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")

    if not records:
        wb.save(EXCEL_DATA_PATH)
        return EXCEL_DATA_PATH

    headers = list(records[0].keys())

    # Write headers
    ws.row_dimensions[1].height = 26
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # Write data rows
    for row_idx, record in enumerate(records, start=2):
        ws.row_dimensions[row_idx].height = 20
        is_even = (row_idx % 2 == 0)
        for col_idx, header in enumerate(headers, start=1):
            val = record.get(header)
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = data_alignment
            if is_even:
                cell.fill = alt_fill

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(EXCEL_DATA_PATH)
    print(f"Excel test data created successfully at: {EXCEL_DATA_PATH}")
    return EXCEL_DATA_PATH


if __name__ == "__main__":
    generate_excel_data()
